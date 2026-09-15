from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Border, Side, Font, Alignment, PatternFill

THIN_BORDER=Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)


def save_conflicts(conflicts, input_filename, date_str, session_str, output_folder):
    conflicts_folder =os.path.join(output_folder, "Conflicts")
    os.makedirs(conflicts_folder, exist_ok=True)
    base_name=os.path.splitext(input_filename)[0]
    conflict_filename=f"Conflicts_{base_name}.xlsx"
    conflict_filepath=os.path.join(conflicts_folder, conflict_filename)
    
    wb=Workbook()
    ws=wb.active
    ws.title="Conflicts"
    ws.append([f"CONFLICTS DETECTED - {date_str} {session_str}"])
    ws['A1'].font =Font(name='Calibri', bold=True, size=14, color='CC0000')
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A1:E1')
    ws.append([])
    
    headers=["USN", "Student Name", "Branch", "Conflicting Subjects", "Subject Count"]
    ws.append(headers)
    for col_idx in range(1, 6):
        cell =ws.cell(row=3, column=col_idx)
        cell.font =Font(name='Calibri', bold=True, size=11)
        cell.fill =PatternFill(start_color='FFE6E6', end_color='FFE6E6', fill_type='solid')
        cell.alignment =Alignment(horizontal='center', vertical='center')
        cell.border=THIN_BORDER
    for idx, conflict in enumerate(conflicts):
        row_data=[conflict['USN'], conflict['Name'], conflict['Branch'], conflict['Subjects'], conflict['SubjectCount']]
        ws.append(row_data)
        row_num=4+idx
        for col_idx in range(1, 6):
            cell =ws.cell(row=row_num, column=col_idx)
            cell.font =Font(name='Calibri', size=10)
            cell.border =THIN_BORDER
            cell.alignment =Alignment(vertical='center')
            if col_idx in [1, 3, 5]:
                cell.alignment =Alignment(horizontal='center', vertical='center')
    
    ws.column_dimensions['A'].width=15
    ws.column_dimensions['B'].width=25
    ws.column_dimensions['C'].width=12
    ws.column_dimensions['D'].width=50
    ws.column_dimensions['E'].width=15
    wb.save(conflict_filepath)
    return conflict_filename
