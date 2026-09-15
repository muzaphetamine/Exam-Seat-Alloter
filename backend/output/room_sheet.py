from openpyxl.styles import Border, Side, Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

THIN_BORDER=Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)
COURSE_HEADER_FILLS=["E6F3FF","E6FFE6","FFF0E6","F0E6FF","FFF2CC","E2F0D9","FCE4D6","DDEBF7"]
COURSE_BODY_FILLS=["F8FCFF","F8FFF8","FFFCF8","FCF8FF","FFFDF2","F5FAF2","FFF8F5","F5F9FD"]


def get_course_fill(course, fill_type="header"):
    course=str(course).strip().upper()
    if not course: color = "F5F5F5"
    else:
        color_index=sum(ord(char) for char in course) %len(COURSE_HEADER_FILLS)
        if fill_type=="body":
            color=COURSE_BODY_FILLS[color_index]
        else:
            color=COURSE_HEADER_FILLS[color_index]
    return PatternFill(start_color=color,end_color=color,fill_type="solid")


def create_room_sheet(wb,room_name,room_layout,course_header,room_students):
    rows=len(room_layout)
    cols=len(room_layout[0]) if room_layout else 0
    ws=wb.create_sheet(title=str(room_name)[:31])

    ws.append([f"Room: {room_name}"])
    ws["A1"].font=Font(name="Calibri",bold=True,size=16)
    ws["A1"].alignment=Alignment(horizontal="center")
    if cols: ws.merge_cells(f"A1:{get_column_letter(cols)}1")

    ws.append(course_header)
    header_row=2
    for col_idx in range(cols):
        cell= ws.cell(row=header_row,column=col_idx+1)
        cell.font =Font(name="Calibri",bold=True,size=11)
        cell.alignment =Alignment(horizontal="center",vertical="center")
        course=""
        if course_header[col_idx]:
            course =str(course_header[col_idx]).split("(")[0]
        cell.fill =get_course_fill(course,"header")

        if col_idx==0:
            cell.border =Border(left=Side(style="thick"),top=Side(style="thick"),bottom=Side(style="thin"))
        elif col_idx==cols-1:
            cell.border =Border(right=Side(style="thick"),top=Side(style="thick"),bottom=Side(style="thin"))
        else:
            cell.border =Border(top=Side(style="thick"),bottom=Side(style="thin"))

    layout_start_row=3
    for row_idx, row_data in enumerate(room_layout):
        ws.append(row_data)
        for col_idx in range(cols):
            cell =ws.cell(row=layout_start_row + row_idx, column=col_idx+1)
            cell.font =Font(name="Calibri", size=10)
            cell.alignment =Alignment(horizontal="center", vertical="center")
            course=""
            if course_header[col_idx]:
                course =str(course_header[col_idx]).split("(")[0]
            cell.fill =get_course_fill(course,"body")

            left =(Side(style="thick") if col_idx==0 else Side(style=None))
            right =(Side(style="thick") if col_idx==cols-1 else Side(style=None))
            bottom =(Side(style="thick") if row_idx==rows-1 else Side(style=None))
            cell.border =Border(left=left, right=right, bottom=bottom)

    for col_idx in range(1, cols+1):
        ws.column_dimensions[get_column_letter(col_idx)].width =16

    ws.append([])
    ws.append([])
    if room_students:
        ws.append(["Students in this room:"])
        students_header_row =rows+5
        ws.cell(
            row=students_header_row,
            column=1
        ).font =Font(
            name="Calibri",
            bold=True,
            size=12
        )
        ws.append(["USN","Name","Course","SubjectCode"])
        table_header_row =students_header_row+1
        for col_idx in range(1, 5):
            cell =ws.cell(row=table_header_row, column=col_idx)
            cell.font =Font(name="Calibri", bold=True, size=10)
            cell.fill =PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
            cell.alignment =Alignment(horizontal="center")
            cell.border=THIN_BORDER

        for idx, student in enumerate(room_students):
            ws.append([
                student["USN"],
                student["Name"],
                student["ExtractedCourse"],
                student["SubjectCode"]
            ])
            row_num =(table_header_row+1+idx)
            fill_color =("F8F9FA" if idx%2==0 else "FFFFFF")
            for col_idx in range(1, 5):
                cell =ws.cell(row=row_num, column=col_idx)
                cell.font =Font(name="Calibri", size=9)
                cell.fill =PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
                cell.border=THIN_BORDER
                if col_idx==1:
                    cell.alignment =Alignment(horizontal="center")
    return ws