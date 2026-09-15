from openpyxl.styles import Border, Side, Font, Alignment, PatternFill

THIN_BORDER=Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
HEADER_FILL=PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
TITLE_FONT=Font(bold=True, size=14, color="FFFFFF")
HEADER_FONT=Font(bold=True, color="FFFFFF")
BODY_FONT=Font(size=10)


def create_room_allotment_sheet(wb, allocations, date_obj, start_dt, end_dt):
    allocations =sorted(allocations, key=lambda x: (x["Branch"], x["SubCode"]))
    ws =wb.create_sheet("Room Allocation", 0)
    date_str =date_obj.strftime("%d-%b-%Y")
    session_str =f"{start_dt.strftime('%I:%M%p')} TO {end_dt.strftime('%I:%M%p')}"
    
    ws.append(["EXAM SEAT ALLOCATION"])
    ws.merge_cells("A1:G1")
    ws["A1"].font=TITLE_FONT
    ws["A1"].fill=HEADER_FILL
    ws["A1"].alignment=Alignment(horizontal="center", vertical="center")
    ws["A1"].border=THIN_BORDER
    ws.row_dimensions[1].height=30

    ws.append([f"DATE : {date_str}    SESSION : {session_str}"])
    ws.merge_cells("A2:G2")
    ws["A2"].font =Font(name='Calibri', bold=True, size=11)
    ws["A2"].alignment=Alignment(horizontal="center")
    headers=["Room No", "BRANCH", "Sub code", "Sub Name", "Sem", "ALLOTTED USN'S", "TOTAL"]
    ws.append(headers)
    
    for cell in ws[3]:
        cell.font=HEADER_FONT
        cell.fill=HEADER_FILL
        cell.border=THIN_BORDER
        cell.alignment=Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height=25
    
    for alloc in allocations:
        ws.append([
            alloc["Room"],
            alloc["Branch"],
            alloc["SubCode"],
            alloc["SubName"],
            alloc["Sem"],
            ", ".join(alloc["USNs"]),
            alloc["Total"]
        ])
        
        r=ws.max_row
        for c in range(1,8):
            cell=ws.cell(row=r, column=c)
            cell.font=BODY_FONT
            cell.border=THIN_BORDER
            cell.alignment=Alignment(
                horizontal="center" if c in [1, 2, 3, 5, 7] else "left",
                vertical="top",
                wrap_text=True
            )
            if (r-4)%2==0:
                cell.fill =PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid')
    
    ws.column_dimensions["A"].width=12
    ws.column_dimensions["B"].width=15
    ws.column_dimensions["C"].width=12
    ws.column_dimensions["D"].width=35
    ws.column_dimensions["E"].width=8
    ws.column_dimensions["F"].width=80
    ws.column_dimensions["G"].width=10


def create_unallocated_sheet(wb, unallocated_students):
    if unallocated_students:
        ws_unalloc = wb.create_sheet(title="Unallocated Students")
        ws_unalloc.append(["Unallocated Students"])
        ws_unalloc['A1'].font = Font(name='Calibri', bold=True, size=14, color='CC0000')
        ws_unalloc.merge_cells('A1:D1')
        ws_unalloc['A1'].alignment = Alignment(horizontal='center')
        ws_unalloc.append([])
        
        ws_unalloc.append(["USN", "Name", "Course", "SubjectCode"])
        for col_idx in range(1,5):
            cell=ws_unalloc.cell(row=3, column=col_idx)
            cell.font=Font(name='Calibri', bold=True)
            cell.fill=PatternFill(start_color='FFE6E6', end_color='FFE6E6', fill_type='solid')
            cell.alignment=Alignment(horizontal='center')
            cell.border=THIN_BORDER
        
        for idx, student in enumerate(unallocated_students):
            ws_unalloc.append([student["USN"], student["Name"], student["ExtractedCourse"], student["SubjectCode"]])   
            row_num =4+idx
            for col_idx in range(1,5):
                cell=ws_unalloc.cell(row=row_num, column=col_idx)
                cell.font=Font(name='Calibri', size=9)
                cell.border=THIN_BORDER
                if col_idx==1:
                    cell.alignment = Alignment(horizontal='center')