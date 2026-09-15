from collections import defaultdict
from openpyxl.styles import Border, Side, Font, Alignment, PatternFill
from backend.core.allocator import extract_course_from_usn

THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)


def create_summary_sheet(
    wb,
    session_students,
    allocated_students,
    unallocated_students,
    rooms_df,
    date_obj,
    start_dt=None,
    end_dt=None,
    session_label=None
):
    total_students=len(session_students)
    room_capacity =(rooms_df["Rows"].astype(int) * rooms_df["Cols"].astype(int))
    avg_capacity =room_capacity.mean()
    date_str =date_obj.strftime("%d-%b-%Y")

    if start_dt is not None and end_dt is not None:
        session_str=(f"{start_dt.strftime('%I:%M%p')} TO {end_dt.strftime('%I:%M%p')}")
    elif session_label:
        session_str=str(session_label)
    else:
        session_str=""

    title=f"Exam Session Summary - {date_str}"
    if session_str: title+=f" {session_str}"
    ws_summary =wb.create_sheet(title="Session Summary")
    ws_summary.append([title])
    ws_summary["A1"].font =Font(name="Calibri", bold=True, size=16, color="FFFFFF")
    ws_summary.merge_cells("A1:C1")
    ws_summary["A1"].alignment =Alignment(horizontal="center")
    ws_summary["A1"].fill =PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    ws_summary.append([])

    extra_rooms =(max(0,-(-len(unallocated_students)//int(avg_capacity))) if avg_capacity>0 else 0)
    stats_data =[
        ["Metric", "Value", "Percentage"],
        ["Total Students", total_students, "100%"],
        [
            "Allocated Students",
            len(allocated_students),
            (f"{len(allocated_students) / total_students * 100:.1f}%" if total_students>0 else "0%")
        ],
        [
            "Unallocated Students",
            len(unallocated_students),
            (f"{len(unallocated_students) / total_students * 100:.1f}%" if total_students>0 else "0%")
        ],
        ["", "", ""],
        ["Average Room Capacity", int(avg_capacity),""],
        ["Rooms Used", len(rooms_df), ""],
        ["Extra Rooms Needed", extra_rooms,""],
    ]

    for row_idx, row_data in enumerate(stats_data):
        ws_summary.append(row_data)
        current_row =row_idx+3
        for col_idx in range(1, 4):
            cell =ws_summary.cell(row=current_row, column=col_idx)
            cell.font =Font(name="Calibri", size=10)
            if row_idx==0:
                cell.font =Font(name="Calibri", bold=True, size=11)
                cell.fill =PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
                cell.alignment =Alignment(horizontal="center")
            elif row_data[0]=="":
                continue
            else:
                if col_idx in (2,3):
                    cell.alignment =Alignment(horizontal="center")

            if row_data[0]!="":
                cell.border=THIN_BORDER

    ws_summary.append([])
    ws_summary.append(["Course Distribution"])
    course_header_row =len(stats_data)+5

    ws_summary.cell(
        row=course_header_row,
        column=1
    ).font =Font(
        name="Calibri",
        bold=True,
        size=12
    )

    ws_summary.cell(
        row=course_header_row,
        column=1
    ).fill =PatternFill(
        start_color="E7E6E6",
        end_color="E7E6E6",
        fill_type="solid"
    )

    ws_summary.append(["Course","Students","Percentage"])

    course_table_header =course_header_row+1
    for col_idx in range(1,4):
        cell = ws_summary.cell(row=course_table_header, column=col_idx)
        cell.font =Font(name="Calibri", bold=True, size=10)
        cell.fill =PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid")
        cell.alignment =Alignment(horizontal="center")
        cell.border=THIN_BORDER

    course_count=defaultdict(int)
    for student in session_students.itertuples():
        course = extract_course_from_usn(student.USN)
        course_count[course]+=1
    for idx, (course, count) in enumerate(course_count.items()):
        percentage =(f"{count / total_students * 100:.1f}%" if total_students > 0 else "0%")
        ws_summary.append([course, count, percentage])
        current_row =(course_table_header+1+idx)
        for col_idx in range(1,4):
            cell =ws_summary.cell(row=current_row, column=col_idx)
            cell.font=Font(name="Calibri", size=9)
            if col_idx in (2,3):
                cell.alignment =Alignment(horizontal="center")
            cell.border=THIN_BORDER
            if idx%2==0:
                cell.fill = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")

    ws_summary.column_dimensions["A"].width=20
    ws_summary.column_dimensions["B"].width=15
    ws_summary.column_dimensions["C"].width=15
    return ws_summary