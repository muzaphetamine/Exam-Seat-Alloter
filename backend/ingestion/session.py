import os
import re
from datetime import datetime, timedelta
import pandas as pd
from openpyxl import load_workbook
from backend.config import EXAM_DURATION_HOURS
from .common import clean_columns, missing_columns

SESSION_STUDENT_COLUMNS = {"USN","Name","SubjectCode"}


def extract_course_from_usn(usn):
    usn=str(usn).strip().upper()
    match = re.match(r'^\d[A-Z]{2}\d{2}([A-Z]{2,4})\d{3}$', usn)
    if not match:
        raise ValueError(f"Could not determine course from USN: {usn}")
    return match.group(1)


def validate_session_file(filepath):
    date_str, time_str =extract_metadata(filepath)
    date_obj, start_dt, end_dt =normalize_datetime(date_str, time_str)
    if not date_obj:
        raise ValueError("could not extract a valid Date and Time from the session metadata")

    ext =os.path.splitext(filepath)[1].lower()
    df=pd.read_excel(filepath, skiprows=3, engine="xlrd" if ext==".xls" else None)
    df=clean_columns(df)
    df=df.rename(columns={
        "S No": "SNo",
        "Student Name": "Name",
        "Subject Code": "SubjectCode",
        "Subject Name": "SubjectName",
        "Semester": "Semester"
    })
    missing=missing_columns(df, SESSION_STUDENT_COLUMNS)
    if missing:
        raise ValueError("missing required session table columns: " + ", ".join(missing))

    student_rows=df.dropna(subset=["USN"])
    for usn in student_rows["USN"]:
        extract_course_from_usn(usn)
    return True


def ingest_session_files(filepaths):
    if not filepaths:
        raise ValueError("No session Excel files were supplied.")
    errors=[]
    for filepath in filepaths:
        try: validate_session_file(filepath)
        except Exception as exc: errors.append(f"{os.path.basename(filepath)}: {exc}")
    if errors:
        raise ValueError("Session input error(s):\n- " + "\n- ".join(errors))
    return filepaths


def extract_metadata(filepath):
    ext =os.path.splitext(filepath)[1].lower()
    text_blob="" 
    if ext==".xlsx":
        wb=load_workbook(filepath)
        ws=wb.active
        for r in range(1, 6):
            for c in range(1, 6):
                val=ws.cell(r, c).value
                if isinstance(val, str):
                    text_blob+=" "+val
    else:
        df =pd.read_excel(filepath, nrows=5, header=None, engine="xlrd")
        text_blob =" ".join(df.fillna("").astype(str).values.flatten())
    date_match =re.search(r'\d{4}-\d{2}-\d{2}', text_blob)
    time_match =re.search(r'\d{2}:\d{2}(:\d{2})?', text_blob)
    date_str =date_match.group() if date_match else None
    time_str=time_match.group() if time_match else None
    return date_str, time_str


def normalize_datetime(date_str, time_str):
    if not date_str or not time_str:
        return None, None, None
    try: date_obj= datetime.strptime(date_str, '%Y-%m-%d')
    except ValueError: return None, None, None
    start_time = None
    for fmt in ('%H:%M:%S', '%H:%M'):
        try:
            start_time = datetime.strptime(time_str, fmt)
            break
        except: pass
    if not start_time:
        return None, None, None
    start_dt=datetime.combine(date_obj.date(), start_time.time())
    end_dt=start_dt + timedelta(hours=EXAM_DURATION_HOURS) 
    return date_obj, start_dt, end_dt


def read_students_from_file(filepath):
    ext=os.path.splitext(filepath)[1].lower()
    df=pd.read_excel(filepath, skiprows=3, engine="xlrd" if ext == ".xls" else None)
    df.columns=df.columns.str.strip()
    df=df.rename(columns={
        'S No': 'SNo',
        'USN': 'USN',
        'Student Name': 'Name',
        'Subject Code': 'SubjectCode',
        'Subject Name': 'SubjectName',
        'Semester': 'Semester'
    })

    df = df.dropna(subset=["USN"]).copy()
    df["USN"] =(df["USN"].astype(str).str.strip().str.upper())
    df["Name"] =(df["Name"].fillna("").astype(str).str.strip())
    df["SubjectCode"] =(df["SubjectCode"].fillna("").astype(str).str.strip().str.upper())
    df["Course"] =df["USN"].apply(extract_course_from_usn)

    return df


def read_excel_table_with_header_offset(filepath, skiprows=0):
    ext=os.path.splitext(filepath)[1].lower()
    return pd.read_excel(
        filepath,
        skiprows=skiprows,
        engine="xlrd" if ext == ".xls" else None
    )