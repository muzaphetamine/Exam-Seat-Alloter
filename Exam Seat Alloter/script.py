import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side, Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import os
import re
from collections import defaultdict
from datetime import datetime, timedelta

SESSION_TIME_MAP = {
    'morning': '09:00',
    'afternoon': '14:00'
}
EXAM_DURATION_HOURS = 3
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
TITLE_FONT = Font(bold=True, size=14, color="FFFFFF")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BODY_FONT = Font(size=10)


def extract_course_from_usn(usn):
    match = re.search(r'\d{2}([A-Z]{2,3})\d{3}$', str(usn))
    return match.group(1) if match else 'UNKNOWN'


def extract_metadata(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    text_blob = "" 
    if ext == ".xlsx":
        wb = load_workbook(filepath)
        ws = wb.active
        for r in range(1, 6):
            for c in range(1, 6):
                val = ws.cell(r, c).value
                if isinstance(val, str):
                    text_blob += " " + val
    else:
        df = pd.read_excel(filepath, nrows=5, header=None, engine="xlrd")
        text_blob = " ".join(df.fillna("").astype(str).values.flatten())

    date_match =re.search(r'\d{4}-\d{2}-\d{2}', text_blob)
    time_match =re.search(r'\d{2}:\d{2}(:\d{2})?', text_blob)
    date_str =date_match.group() if date_match else None
    time_str=time_match.group() if time_match else None
    return date_str, time_str


def normalize_datetime(date_str, time_str):
    if not date_str or not time_str:
        return None, None, None
    try:
        date_obj= datetime.strptime(date_str, '%Y-%m-%d')
    except:
        return None, None, None
    start_time = None
    for fmt in ('%H:%M:%S', '%H:%M'):
        try:
            start_time = datetime.strptime(time_str, fmt)
            break
        except:
            pass
    if not start_time:
        return None, None, None
    start_dt= datetime.combine(date_obj.date(), start_time.time())
    end_dt =start_dt + timedelta(hours=EXAM_DURATION_HOURS) 
    return date_obj, start_dt, end_dt


def read_students_from_file(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    df = pd.read_excel(filepath, skiprows=3, engine="xlrd" if ext == ".xls" else None)
    df.columns = df.columns.str.strip()
    df = df.rename(columns={
        'S No': 'SNo',
        'USN': 'USN',
        'Student Name': 'Name',
        'Subject Code': 'SubjectCode',
        'Subject Name': 'SubjectName',
        'Semester': 'Semester'
    })
    df = df.dropna(subset=['USN'])
    return df


def check_conflicts(students_df):
    conflicts = [] 
    for usn, group in students_df.groupby('USN'):
        if len(group) > 1:
            student_data=group.iloc[0]
            conflicts.append({
                'USN': usn,
                'Name': student_data['Name'],
                'Branch': extract_course_from_usn(usn),
                'Subjects': ', '.join(f"{row['SubjectCode']} ({row.get('SubjectName', row['SubjectCode'])})" for _, row in group.iterrows()),
                'SubjectCount': len(group)
            })
    return conflicts


def group_students_by_course_subject(session_students):
    session_students['ExtractedCourse'] = session_students['USN'].apply(extract_course_from_usn)
    course_subject_groups = defaultdict(list)
    for _, student in session_students.iterrows():
        key = (student['ExtractedCourse'], student['SubjectCode'])
        course_subject_groups[key].append(student)  
    return dict(course_subject_groups)


def create_room_pairs(course_subject_groups):
    groups = list(course_subject_groups.keys())
    room_pairs = []
    used_groups = set()
    for i, group1 in enumerate(groups):
        if group1 in used_groups:
            continue  
        course1, subject1 = group1
        best_pair = None
        for j, group2 in enumerate(groups[i+1:], i+1):
            if group2 in used_groups:
                continue
            course2, subject2 = group2
            if course1 != course2 and subject1 != subject2:
                best_pair = group2
                break
        
        if not best_pair:
            for j, group2 in enumerate(groups[i+1:], i+1):
                if group2 in used_groups:
                    continue     
                course2, subject2 = group2
                if course1 != course2:
                    best_pair = group2
                    break

        if not best_pair:
            for j, group2 in enumerate(groups[i+1:], i+1):
                if group2 in used_groups:
                    continue
                best_pair = group2
                break
        
        if best_pair:
            room_pairs.append((group1, best_pair))
            used_groups.add(group1)
            used_groups.add(best_pair)
        else:
            room_pairs.append((group1, None))
            used_groups.add(group1)
    return room_pairs


def allocate_students_to_room(group1_students, group2_students, rows, cols):
    room_layout = [['' for _ in range(cols)] for _ in range(rows)]
    course_header = [''] * cols
    room_students = []
    group1_idx = 0
    group2_idx = 0
    
    group1_course =group1_students[0]['ExtractedCourse'] if group1_students else ''
    group1_subject =group1_students[0]['SubjectCode'] if group1_students else ''
    group2_course =group2_students[0]['ExtractedCourse'] if group2_students else ''
    group2_subject =group2_students[0]['SubjectCode'] if group2_students else ''
    
    for col in range(cols):
        if col % 2 == 0:  # Even columns for group 1
            current_group = group1_students
            current_idx = group1_idx
            course_header[col] = f"{group1_course}({group1_subject})" if group1_course else ''
        else:  # Odd columns for group 2
            current_group = group2_students if group2_students else group1_students
            current_idx = group2_idx if group2_students else group1_idx
            course_header[col] = f"{group2_course}({group2_subject})" if group2_students and group2_course else f"{group1_course}({group1_subject})" if group1_course else ''
        
        for row in range(rows):
            if current_idx < len(current_group):
                student = current_group[current_idx]
                room_layout[row][col] = student['USN']
                room_students.append(student)
                current_idx += 1
            else:
                break
        
        if col % 2 == 0 or not group2_students:
            group1_idx = current_idx
        else:
            group2_idx = current_idx
    return room_layout, room_students, course_header

def create_room_allotment_sheet(wb, allocations, date_obj, start_dt, end_dt):
    allocations = sorted(
        allocations,
        key=lambda x: (x["Branch"], x["SubCode"])
    )
    ws = wb.create_sheet("Room Allocation", 0)
    date_str = date_obj.strftime("%d-%b-%Y")
    session_str = f"{start_dt.strftime('%I:%M%p')} TO {end_dt.strftime('%I:%M%p')}"
    
    ws.append(["VTU Theory Exam - CANDIDATE EXAM HALL ALLOTMENT"])
    ws.merge_cells("A1:G1")
    ws["A1"].font = TITLE_FONT
    ws["A1"].fill = HEADER_FILL
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws["A1"].border = THIN_BORDER
    ws.row_dimensions[1].height = 30

    ws.append([f"DATE : {date_str}    SESSION : {session_str}"])
    ws.merge_cells("A2:G2")
    ws["A2"].font = Font(name='Calibri', bold=True, size=11)
    ws["A2"].alignment = Alignment(horizontal="center")
    headers = ["Room No", "BRANCH", "Sub code", "Sub Name", "Sem", "ALLOTTED USN'S", "TOTAL"]
    ws.append(headers)
    
    for cell in ws[3]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 25
    
    for alloc in allocations:
        ws.append([
            alloc["Room"],
            alloc["Branch"],
            alloc["SubCode"],
            alloc["SubName"],
            alloc["Sem"],
            ", ".join(alloc["USNs"]),
            alloc["Total"]
        ])
        
        r = ws.max_row
        for c in range(1, 8):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.border = THIN_BORDER
            cell.alignment = Alignment(
                horizontal="center" if c in [1, 2, 3, 5, 7] else "left",
                vertical="top",
                wrap_text=True
            )
            if (r - 4) % 2 == 0:
                cell.fill = PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid')
    
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 35
    ws.column_dimensions["E"].width = 8
    ws.column_dimensions["F"].width = 80
    ws.column_dimensions["G"].width = 10

def create_output_file(session_students, rooms_df, date_obj, start_dt, end_dt, output_folder):
    course_subject_groups = group_students_by_course_subject(session_students)
    room_pairs = create_room_pairs(course_subject_groups)
    wb = Workbook()
    if wb.active:
        wb.remove(wb.active)
    all_allocated_students = []
    room_allocations = []
    room_pair_idx = 0

    for room_idx, room_row in rooms_df.iterrows():
        room_name = room_row["RoomName"]
        rows = int(room_row["Rows"])
        cols = int(room_row["Cols"])
        if room_pair_idx < len(room_pairs):
            pair = room_pairs[room_pair_idx]
            group1_key, group2_key = pair
            group1_students= course_subject_groups[group1_key]
            group2_students= course_subject_groups[group2_key] if group2_key else []

            room_layout, room_students, course_header = allocate_students_to_room(
                group1_students, group2_students, rows, cols
            )

            for student in room_students:
                branch = student['ExtractedCourse']
                subcode = student['SubjectCode']
                found = False
                for alloc in room_allocations:
                    if (alloc['Room'] == room_name and 
                        alloc['Branch'] == branch and 
                        alloc['SubCode'] == subcode):
                        alloc['USNs'].append(student['USN'])
                        alloc['Total'] += 1
                        found = True
                        break
            
                if not found:
                    room_allocations.append({
                        'Room': room_name,
                        'Branch': branch,
                        'SubCode': subcode,
                        'SubName': student.get('SubjectName', student['SubjectCode']),
                        'Sem': student.get('Semester', ''),
                        'USNs': [student['USN']],
                        'Total': 1
                    })

            course_subject_groups[group1_key] = course_subject_groups[group1_key][len([s for s in room_students if s['ExtractedCourse'] == group1_key[0]]):]
            if group2_key:
                course_subject_groups[group2_key] = course_subject_groups[group2_key][len([s for s in room_students if s['ExtractedCourse'] == group2_key[0]]):]

            if len(course_subject_groups[group1_key]) == 0:
                if group2_key and len(course_subject_groups[group2_key]) == 0:
                    room_pair_idx += 1
                elif not group2_key:
                    room_pair_idx += 1
            all_allocated_students.extend(room_students)

        else:
            room_layout = [['' for _ in range(cols)] for _ in range(rows)]
            course_header = [''] * cols
            room_students = []

        if not room_students:
            continue
        ws = wb.create_sheet(title=room_name[:31])
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(name='Calibri')
        ws.append([f"Room: {room_name}"])
        ws['A1'].font = Font(name='Calibri', bold=True, size=16)
        ws['A1'].alignment = Alignment(horizontal='center')
        ws.merge_cells(f'A1:{get_column_letter(cols)}1')
        
        ws.append(course_header)
        header_row = 2
        for col_idx in range(cols):
            cell = ws.cell(row=header_row, column=col_idx + 1)
            cell.font = Font(name='Calibri', bold=True, size=11)
            cell.alignment = Alignment(horizontal='center', vertical='center')

            if 'CS' in str(cell.value):
                cell.fill = PatternFill(start_color='E6F3FF', end_color='E6F3FF', fill_type='solid')
            elif 'AI' in str(cell.value):
                cell.fill = PatternFill(start_color='E6FFE6', end_color='E6FFE6', fill_type='solid')
            elif 'DS' in str(cell.value):
                cell.fill = PatternFill(start_color='FFF0E6', end_color='FFF0E6', fill_type='solid')
            elif 'ML' in str(cell.value):
                cell.fill = PatternFill(start_color='F0E6FF', end_color='F0E6FF', fill_type='solid')
            else:
                cell.fill = PatternFill(start_color='F5F5F5', end_color='F5F5F5', fill_type='solid')
            
            if col_idx == 0:
                cell.border = Border(left=Side(style='thick'), top=Side(style='thick'), bottom=Side(style='thin'))
            elif col_idx == cols - 1:
                cell.border = Border(right=Side(style='thick'), top=Side(style='thick'), bottom=Side(style='thin'))
            else:
                cell.border = Border(top=Side(style='thick'), bottom=Side(style='thin'))

        layout_start_row = 3
        for row_idx, row_data in enumerate(room_layout):
            ws.append(row_data)
            for col_idx in range(cols):
                cell = ws.cell(row=layout_start_row + row_idx, column=col_idx + 1)
                cell.font = Font(name='Calibri', size=10)
                cell.alignment = Alignment(horizontal='center', vertical='center')
                if col_idx % 2 == 0:
                    if 'CS' in course_header[col_idx]:
                        cell.fill = PatternFill(start_color='F8FCFF', end_color='F8FCFF', fill_type='solid')
                    elif 'AI' in course_header[col_idx]:
                        cell.fill = PatternFill(start_color='F8FFF8', end_color='F8FFF8', fill_type='solid')
                    elif 'DS' in course_header[col_idx]:
                        cell.fill = PatternFill(start_color='FFFCF8', end_color='FFFCF8', fill_type='solid')
                    elif 'ML' in course_header[col_idx]:
                        cell.fill = PatternFill(start_color='FCF8FF', end_color='FCF8FF', fill_type='solid')
                    else:
                        cell.fill = PatternFill(start_color='FAFAFA', end_color='FAFAFA', fill_type='solid')

                border = Border()
                if col_idx == 0:
                    border = Border(left=Side(style='thick'))
                if col_idx == cols - 1:
                    border = Border(right=Side(style='thick')) if col_idx == 0 else Border(left=border.left, right=Side(style='thick'))
                if row_idx == len(room_layout) - 1:
                    if col_idx == 0:
                        border = Border(left=Side(style='thick'), bottom=Side(style='thick'))
                    elif col_idx == cols - 1:
                        border = Border(right=Side(style='thick'), bottom=Side(style='thick'))
                    else:
                        border = Border(bottom=Side(style='thick'))
                cell.border = border

        for col_idx in range(1, cols + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 16
        ws.append([])
        ws.append([])

        if room_students:
            ws.append(["Students in this room:"])
            students_header_row = len(room_layout) + 5
            ws.cell(row=students_header_row, column=1).font = Font(name='Calibri', bold=True, size=12)
            ws.append(["USN", "Name", "Course", "SubjectCode"])
            table_header_row = students_header_row + 1
            for col_idx in range(1, 5):
                cell = ws.cell(row=table_header_row, column=col_idx)
                cell.font = Font(name='Calibri', bold=True, size=10)
                cell.fill = PatternFill(start_color='D9EAF7', end_color='D9EAF7', fill_type='solid')
                cell.alignment = Alignment(horizontal='center')
                cell.border = THIN_BORDER

            for idx, student in enumerate(room_students):
                ws.append([student["USN"], student["Name"], student["ExtractedCourse"], student["SubjectCode"]])  
                row_num = table_header_row + 1 + idx
                fill_color = 'F8F9FA' if idx % 2 == 0 else 'FFFFFF'
                for col_idx in range(1, 5):
                    cell = ws.cell(row=row_num, column=col_idx)
                    cell.font = Font(name='Calibri', size=9)
                    cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')
                    cell.border = THIN_BORDER
                    if col_idx == 1:
                        cell.alignment = Alignment(horizontal='center')
    
    unallocated_students = []
    for group_key, remaining_students in course_subject_groups.items():
        unallocated_students.extend(remaining_students)
    
    create_room_allotment_sheet(wb, room_allocations, date_obj, start_dt, end_dt)
    
    # Unallocated sheet
    if unallocated_students:
        ws_unalloc = wb.create_sheet(title="Unallocated Students")
        
        ws_unalloc.append(["Unallocated Students"])
        ws_unalloc['A1'].font = Font(name='Calibri', bold=True, size=14, color='CC0000')
        ws_unalloc.merge_cells('A1:D1')
        ws_unalloc['A1'].alignment = Alignment(horizontal='center')
        ws_unalloc.append([])
        
        ws_unalloc.append(["USN", "Name", "Course", "SubjectCode"])
        for col_idx in range(1, 5):
            cell = ws_unalloc.cell(row=3, column=col_idx)
            cell.font = Font(name='Calibri', bold=True)
            cell.fill = PatternFill(start_color='FFE6E6', end_color='FFE6E6', fill_type='solid')
            cell.alignment = Alignment(horizontal='center')
            cell.border = THIN_BORDER
        
        for idx, student in enumerate(unallocated_students):
            ws_unalloc.append([student["USN"], student["Name"], student["ExtractedCourse"], student["SubjectCode"]])   
            row_num = 4 + idx
            for col_idx in range(1, 5):
                cell = ws_unalloc.cell(row=row_num, column=col_idx)
                cell.font = Font(name='Calibri', size=9)
                cell.border = THIN_BORDER
                if col_idx == 1:
                    cell.alignment = Alignment(horizontal='center')
    
    total_students = len(session_students)
    average_room_capacity = rooms_df['Rows'].astype(int) * rooms_df['Cols'].astype(int)
    avg_capacity = average_room_capacity.mean()
    ws_summary = wb.create_sheet(title="Session Summary")
    date_str = date_obj.strftime("%d-%b-%Y")
    session_str = f"{start_dt.strftime('%I:%M%p')} TO {end_dt.strftime('%I:%M%p')}"
    ws_summary.append([f"Exam Session Summary - {date_str} {session_str}"])
    ws_summary['A1'].font = Font(name='Calibri', bold=True, size=16, color='FFFFFF')
    ws_summary.merge_cells('A1:C1')
    ws_summary['A1'].alignment = Alignment(horizontal='center')
    ws_summary['A1'].fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    ws_summary.append([])
    
    stats_data = [
        ["Metric", "Value", "Percentage"],
        ["Total Students", total_students, "100%"],
        ["Allocated Students", len(all_allocated_students), f"{len(all_allocated_students)/total_students*100:.1f}%" if total_students > 0 else "0%"],
        ["Unallocated Students", len(unallocated_students), f"{len(unallocated_students)/total_students*100:.1f}%" if total_students > 0 else "0%"],
        ["", "", ""],
        ["Average Room Capacity", int(avg_capacity), ""],
        ["Rooms Used", len(rooms_df), ""],
    ]
    
    for row_idx, row_data in enumerate(stats_data):
        ws_summary.append(row_data)
        current_row = row_idx + 3
        for col_idx in range(1, 4):
            cell = ws_summary.cell(row=current_row, column=col_idx)
            cell.font = Font(name='Calibri', size=10) 
            if row_idx == 0:
                cell.font = Font(name='Calibri', bold=True, size=11)
                cell.fill = PatternFill(start_color='D9EAF7', end_color='D9EAF7', fill_type='solid')
                cell.alignment = Alignment(horizontal='center')
            elif row_data[0] == "":
                continue
            else:
                if col_idx ==2 or col_idx==3:
                    cell.alignment = Alignment(horizontal='center')
            if row_data[0] != "":
                cell.border = THIN_BORDER

    ws_summary.append([])
    ws_summary.append(["Course Distribution"])
    course_header_row = len(stats_data) + 5
    ws_summary.cell(row=course_header_row, column=1).font = Font(name='Calibri', bold=True, size=12)
    ws_summary.cell(row=course_header_row, column=1).fill = PatternFill(start_color='E7E6E6', end_color='E7E6E6', fill_type='solid')
    ws_summary.append(["Course", "Students", "Percentage"])
    course_table_header = course_header_row + 1
    for col_idx in range(1, 4):
        cell = ws_summary.cell(row=course_table_header, column=col_idx)
        cell.font = Font(name='Calibri', bold=True, size=10)
        cell.fill = PatternFill(start_color='D9EAF7', end_color='D9EAF7', fill_type='solid')
        cell.alignment = Alignment(horizontal='center')
        cell.border = THIN_BORDER
    
    course_count = defaultdict(int)
    for student in session_students.itertuples():
        course = extract_course_from_usn(student.USN)
        course_count[course] += 1
    for idx, (course, count) in enumerate(course_count.items()):
        percentage = f"{count/total_students*100:.1f}%" if total_students > 0 else "0%"
        ws_summary.append([course, count, percentage])
        current_row = course_table_header + 1 + idx
        for col_idx in range(1, 4):
            cell = ws_summary.cell(row=current_row, column=col_idx)
            cell.font = Font(name='Calibri', size=9)
            if col_idx == 2 or col_idx == 3:
                cell.alignment = Alignment(horizontal='center')
            cell.border = THIN_BORDER
            if idx % 2 == 0:
                cell.fill = PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid')
    
    ws_summary.column_dimensions['A'].width = 20
    ws_summary.column_dimensions['B'].width = 15
    ws_summary.column_dimensions['C'].width = 15
    
    date_str = date_obj.strftime("%d-%b-%Y")
    session_str = f"{start_dt.strftime('%I-%M%p')} TO {end_dt.strftime('%I-%M%p')}"
    filename = f"Roomallotment on {date_str} {session_str}.xlsx"
    filepath = os.path.join(output_folder, filename)
    wb.save(filepath)
    return filename, len(all_allocated_students), len(unallocated_students)


def save_conflicts(conflicts, input_filename, date_str, session_str, output_folder):
    conflicts_folder = os.path.join(output_folder, "Conflicts")
    os.makedirs(conflicts_folder, exist_ok=True)
    base_name = os.path.splitext(input_filename)[0]
    conflict_filename = f"Conflicts_{base_name}.xlsx"
    conflict_filepath = os.path.join(conflicts_folder, conflict_filename)
    wb = Workbook()
    ws = wb.active
    ws.title = "Conflicts"
    ws.append([f"CONFLICTS DETECTED - {date_str} {session_str}"])
    ws['A1'].font = Font(name='Calibri', bold=True, size=14, color='CC0000')
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A1:E1')
    ws.append([])
    
    headers = ["USN", "Student Name", "Branch", "Conflicting Subjects", "Subject Count"]
    ws.append(headers)
    for col_idx in range(1, 6):
        cell = ws.cell(row=3, column=col_idx)
        cell.font = Font(name='Calibri', bold=True, size=11)
        cell.fill = PatternFill(start_color='FFE6E6', end_color='FFE6E6', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = THIN_BORDER
    for idx, conflict in enumerate(conflicts):
        row_data = [conflict['USN'], conflict['Name'], conflict['Branch'], conflict['Subjects'], conflict['SubjectCount']]
        ws.append(row_data)
        row_num = 4 + idx
        for col_idx in range(1, 6):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.font = Font(name='Calibri', size=10)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical='center')
            if col_idx in [1, 3, 5]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
    
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 25
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 50
    ws.column_dimensions['E'].width = 15
    wb.save(conflict_filepath)
    return conflict_filename


def process_session_files(session_files, rooms_df, output_folder):
    results=[]
    
    for session_file in session_files:
        try:
            date_str, time_str = extract_metadata(session_file)
            date_obj, start_dt, end_dt = normalize_datetime(date_str, time_str)
            if not date_obj:
                results.append({
                    'file': os.path.basename(session_file),
                    'status': 'error',
                    'message': 'Could not extract date/time from file'
                })
                continue
            students_df = read_students_from_file(session_file)

            conflicts = check_conflicts(students_df)
            if conflicts:
                conflict_file = save_conflicts(
                    conflicts, 
                    os.path.basename(session_file), 
                    date_str, 
                    time_str, 
                    output_folder
                )
                results.append({
                    'file': os.path.basename(session_file),
                    'status': 'conflict',
                    'message': f'{len(conflicts)} conflicts found',
                    'conflict_file': conflict_file
                })
                continue

            filename, allocated, unallocated = create_output_file(
                students_df, 
                rooms_df, 
                date_obj, 
                start_dt, 
                end_dt, 
                output_folder
            )
            results.append({
                'file': os.path.basename(session_file),
                'status': 'success',
                'output': filename,
                'allocated': allocated,
                'unallocated': unallocated,
                'total': len(students_df)
            })
            
        except Exception as e:
            results.append({
                'file': os.path.basename(session_file),
                'status': 'error',
                'message': str(e)
            })
    return results

def generate_sessions_from_centralized(students_df, schedule_df, temp_input_dir):
    session_files=[]

    for _, sched in schedule_df.iterrows():
        raw_codes = str(sched['SubjectCode'])
        subject_codes = [
            code.strip()
            for code in re.split(r'[;,/]', raw_codes)
            if code.strip()
        ]
        session_students = students_df[
            students_df['SubjectCode'].isin(subject_codes)
        ]
        if session_students.empty:
            continue
        date_val = sched['Date']
        if hasattr(date_val, 'strftime'):
            date_str = date_val.strftime('%Y-%m-%d')
        else:
            date_str = str(date_val).replace('/', '-')
        session_str = str(sched['Session']).replace(' ', '_')
        time_str = SESSION_TIME_MAP.get(session_str.lower())
        session_filename =f"{date_str}_{session_str}.xlsx"
        session_path = os.path.join(temp_input_dir, session_filename)
        with pd.ExcelWriter(session_path, engine='openpyxl') as writer:
            meta_rows = [f"Date: {date_str}"]
            if time_str: meta_rows.append(f"Time: {time_str}")
            else: meta_rows.append(f"Session: {session_str}")
            pd.DataFrame({'A': meta_rows}).to_excel(
                writer, index=False, header=False
            )
            session_students.to_excel(
                writer, startrow=3, index=False
            )
        session_files.append(session_path)
    return session_files


def main_cli():
    input_folder = "Input"
    output_folder = "Output"
    os.makedirs(input_folder, exist_ok=True)
    os.makedirs(output_folder, exist_ok=True)
    if not os.path.exists('Rooms.xlsx'):
        print("ERROR: Rooms.xlsx not found!")
        return
    rooms_df = pd.read_excel('Rooms.xlsx') 
    session_files = [
        os.path.join(input_folder, f) 
        for f in os.listdir(input_folder) 
        if f.endswith(('.xls', '.xlsx'))
    ]
    if not session_files:
        print("ERROR: No session files found in Input folder!")
        return
    print(f"Found {len(session_files)} session file(s)")
    print(f"Loaded {len(rooms_df)} rooms\n")
    
    results = process_session_files(session_files, rooms_df, output_folder)
    print("\n" + "="*60)
    print("PROCESSING SUMMARY")
    print("="*60)
    success_count = sum(1 for r in results if r['status'] == 'success')
    conflict_count = sum(1 for r in results if r['status'] == 'conflict')
    error_count = sum(1 for r in results if r['status'] == 'error')
    print(f"Total files: {len(results)}")
    print(f"Successful: {success_count}")
    print(f"Conflicts: {conflict_count}")
    print(f"Errors: {error_count}")
    print("="*60)
    
    for result in results:
        if result['status'] == 'success':
            print(f"\n{result['file']}")
            print(f"  Output: {result['output']}")
            print(f"  Allocated: {result['allocated']}/{result['total']}")
            if result['unallocated'] > 0:
                print(f"  WARNING: {result['unallocated']} students unallocated")

if __name__ == "__main__":
    main_cli()